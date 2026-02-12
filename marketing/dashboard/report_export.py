"""
Weekly Report Word document export using python-docx.
"""

import io

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT


def _fmt_pct(val):
    if val is None:
        return "—"
    return f"{'+' if val >= 0 else ''}{val:.1f}%"


def _fmt_dollar(val):
    if val is None or val == 0:
        return "—"
    return f"${val:,.0f}"


def _fmt_mts(val):
    if val is None:
        return "—"
    return f"{val * 100:.1f}%"


def _set_cell_text(cell, text, bold=False, size=9):
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(str(text))
    run.font.size = Pt(size)
    if bold:
        run.bold = True


def build_weekly_report_docx(vertical_name, week_start, week_end, snapshot, brand_rows, report):
    """Build a Word document for the weekly report. Returns bytes."""
    doc = Document()

    # ── Title ────────────────────────────────────────────────────
    title = doc.add_heading(f"Weekly Report — {vertical_name}", level=1)
    date_range = f"{week_start.strftime('%b %d')}–{week_end.strftime('%b %d, %Y')}"
    doc.add_paragraph(date_range, style="Subtitle")

    # ── Executive Snapshot ───────────────────────────────────────
    doc.add_heading("Executive Snapshot", level=2)

    table = doc.add_table(rows=5, cols=2, style="Light Grid Accent 1")
    table.alignment = WD_TABLE_ALIGNMENT.LEFT

    metrics = [
        ("Revenue", f"{_fmt_dollar(snapshot['revenue'])}  (WoW {_fmt_pct(snapshot['rev_wow_pct'])}, YoY {_fmt_pct(snapshot['rev_yoy_pct'])})"),
        ("Spend", f"{_fmt_dollar(snapshot['spend'])}  (WoW {_fmt_pct(snapshot['spend_wow_pct'])}, YoY {_fmt_pct(snapshot['spend_yoy_pct'])})"),
        ("MTS", f"{_fmt_mts(snapshot['mts'])}  (goal: {_fmt_mts(snapshot['mts_goal'])})"),
        ("Top Revenue Driver", snapshot["top_driver"]),
        ("Largest YoY Shift", snapshot["largest_yoy_shift"]),
    ]

    for i, (label, value) in enumerate(metrics):
        _set_cell_text(table.cell(i, 0), label, bold=True)
        _set_cell_text(table.cell(i, 1), value)

    # ── Performance by MFG ───────────────────────────────────────
    doc.add_heading("Performance by MFG", level=2)

    cols = ["Brand", "Rev WoW%", "Rev YoY%", "Spend YoY%", "MTS", "Status", "Notes"]
    t2 = doc.add_table(rows=1 + len(brand_rows), cols=len(cols), style="Light Grid Accent 1")
    t2.alignment = WD_TABLE_ALIGNMENT.LEFT

    for j, col_name in enumerate(cols):
        _set_cell_text(t2.cell(0, j), col_name, bold=True, size=8)

    brand_notes = report.brand_notes if report else {}
    for i, row in enumerate(brand_rows, start=1):
        _set_cell_text(t2.cell(i, 0), row["name"], bold=True, size=8)
        _set_cell_text(t2.cell(i, 1), _fmt_pct(row["rev_wow_pct"]), size=8)
        _set_cell_text(t2.cell(i, 2), _fmt_pct(row["rev_yoy_pct"]), size=8)
        _set_cell_text(t2.cell(i, 3), _fmt_pct(row["spend_yoy_pct"]), size=8)
        _set_cell_text(t2.cell(i, 4), _fmt_mts(row["mts"]), size=8)
        status = ", ".join(b["label"] for b in row["alert_badges"]) or "—"
        _set_cell_text(t2.cell(i, 5), status, size=8)
        note = brand_notes.get(str(row["id"]), "")
        _set_cell_text(t2.cell(i, 6), note, size=8)

    # ── Commentary sections ──────────────────────────────────────
    sections = [
        ("Summary Statement", report.summary_statement if report else ""),
        ("Major YoY Shifts", report.major_yoy_shifts if report else ""),
        ("What's Working Well", report.whats_working_well if report else ""),
        ("What Needs Attention", report.whats_needs_attention if report else ""),
        ("What We're Doing About It", report.what_were_doing if report else ""),
        ("Platform Testing & Experiments", report.platform_testing if report else ""),
        ("Channel Mix Observations", report.channel_mix_observations if report else ""),
        ("Risk & Opportunity Outlook (Next 2-4 Weeks)", report.risk_opportunity_outlook if report else ""),
        ("GM Discussion Points", report.gm_discussion_points if report else ""),
    ]

    for heading, content in sections:
        doc.add_heading(heading, level=2)
        doc.add_paragraph(content or "(No content entered)")

    # ── Write to bytes ───────────────────────────────────────────
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.getvalue()
