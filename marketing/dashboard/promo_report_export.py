"""
Promotional Report Word document export using python-docx.
"""

import io

from docx import Document
from docx.shared import Pt
from docx.enum.table import WD_TABLE_ALIGNMENT

from .report_export import _fmt_dollar, _fmt_mts, _fmt_pct, _set_cell_text


def _add_snapshot_table(doc, title, snapshot):
    """Add an executive snapshot table for one time window."""
    doc.add_heading(title, level=2)

    table = doc.add_table(rows=5, cols=2, style="Light Grid Accent 1")
    table.alignment = WD_TABLE_ALIGNMENT.LEFT

    metrics = [
        ("Revenue", f"{_fmt_dollar(snapshot['revenue'])}  (YoY {_fmt_pct(snapshot['rev_yoy_pct'])})"),
        ("Spend", f"{_fmt_dollar(snapshot['spend'])}  (YoY {_fmt_pct(snapshot['spend_yoy_pct'])})"),
        ("MTS", f"{_fmt_mts(snapshot['mts'])}  (goal: {_fmt_mts(snapshot['mts_goal'])})"),
        ("Top Revenue Driver", snapshot["top_driver"]),
        ("Largest YoY Shift", snapshot["largest_yoy_shift"]),
    ]

    for i, (label, value) in enumerate(metrics):
        _set_cell_text(table.cell(i, 0), label, bold=True)
        _set_cell_text(table.cell(i, 1), value)


def build_promo_report_docx(
    vertical_name, promo_label, promo_start, promo_end,
    cmp_start, cmp_end, snapshot_ptd, snapshot_yesterday,
    brand_rows, report,
):
    """Build a Word document for the promotional report. Returns bytes."""
    doc = Document()

    # ── Title ────────────────────────────────────────────────────
    doc.add_heading(f"Promotional Report — {vertical_name} — {promo_label}", level=1)
    promo_range = f"{promo_start.strftime('%b %d')}–{promo_end.strftime('%b %d, %Y')}"
    cmp_range = f"{cmp_start.strftime('%b %d')}–{cmp_end.strftime('%b %d, %Y')}"
    doc.add_paragraph(f"Promo: {promo_range}  vs  Comparison: {cmp_range}", style="Subtitle")

    # ── Executive Snapshots ───────────────────────────────────────
    if snapshot_ptd:
        _add_snapshot_table(doc, "Promo Period-to-Yesterday", snapshot_ptd)

    if snapshot_yesterday:
        _add_snapshot_table(doc, "Yesterday Only", snapshot_yesterday)

    # ── Performance by MFG ───────────────────────────────────────
    doc.add_heading("Performance by MFG", level=2)

    cols = ["Brand", "Rev YoY%", "Spend YoY%", "MTS", "Status", "Notes"]
    t2 = doc.add_table(rows=1 + len(brand_rows), cols=len(cols), style="Light Grid Accent 1")
    t2.alignment = WD_TABLE_ALIGNMENT.LEFT

    for j, col_name in enumerate(cols):
        _set_cell_text(t2.cell(0, j), col_name, bold=True, size=8)

    brand_notes = report.brand_notes if report else {}
    for i, row in enumerate(brand_rows, start=1):
        _set_cell_text(t2.cell(i, 0), row["name"], bold=True, size=8)
        _set_cell_text(t2.cell(i, 1), _fmt_pct(row["rev_yoy_pct"]), size=8)
        _set_cell_text(t2.cell(i, 2), _fmt_pct(row["spend_yoy_pct"]), size=8)
        _set_cell_text(t2.cell(i, 3), _fmt_mts(row["mts"]), size=8)
        status = ", ".join(b["label"] for b in row["alert_badges"]) or "—"
        _set_cell_text(t2.cell(i, 4), status, size=8)
        note = brand_notes.get(str(row["id"]), "")
        _set_cell_text(t2.cell(i, 5), note, size=8)

    # ── Commentary sections ──────────────────────────────────────
    sections = [
        ("Summary Statement", report.summary_statement if report else ""),
        ("Major YoY Shifts", report.major_yoy_shifts if report else ""),
        ("What's Working Well", report.whats_working_well if report else ""),
        ("What Needs Attention", report.whats_needs_attention if report else ""),
        ("What We're Doing About It", report.what_were_doing if report else ""),
        ("GM Discussion Points", report.gm_discussion_points if report else ""),
    ]

    for heading, content in sections:
        doc.add_heading(heading, level=2)
        doc.add_paragraph(content or "(No content entered)")

    # ── Write to bytes ───────────────────────────────────────────
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.getvalue()
